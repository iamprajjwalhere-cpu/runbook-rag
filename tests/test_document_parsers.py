"""Parser checks for text, PDF page locations, and PowerPoint slide structure."""

from io import BytesIO
import sys
import types
import unittest
from unittest.mock import patch

from pptx import Presentation
from pptx.util import Inches

from workspace_rag import _document_chunks, _extract_sections


class DocumentParserTests(unittest.TestCase):
    def test_text_file_extracts_text_without_location_metadata(self):
        sections, warnings = _extract_sections("policy.txt", b"Annual leave is 20 days.")

        self.assertEqual(warnings, [])
        self.assertEqual(sections[0]["text"], "Annual leave is 20 days.")
        self.assertNotIn("page_number", sections[0])

    def test_pdf_extraction_keeps_page_numbers_and_flags_scanned_pages(self):
        pypdf_stub = types.ModuleType("pypdf")

        class FakePage:
            def __init__(self, text):
                self.text = text

            def extract_text(self):
                return self.text

        class FakeReader:
            def __init__(self, _stream, strict=False):
                self.is_encrypted = False
                self.pages = [FakePage("Safety policy text"), FakePage("")]

        pypdf_stub.PdfReader = FakeReader
        with patch.dict(sys.modules, {"pypdf": pypdf_stub}):
            sections, warnings = _extract_sections("safety.pdf", b"mock pdf")

        self.assertEqual(sections[0]["page_number"], 1)
        self.assertEqual(sections[0]["section"], "Page 1")
        self.assertIn("Page 2", warnings[0])

    def test_pdf_chunks_keep_page_metadata(self):
        pypdf_stub = types.ModuleType("pypdf")

        class FakePage:
            def extract_text(self):
                return "A page-specific operating procedure."

        class FakeReader:
            def __init__(self, _stream, strict=False):
                self.is_encrypted = False
                self.pages = [FakePage()]

        pypdf_stub.PdfReader = FakeReader
        with patch.dict(sys.modules, {"pypdf": pypdf_stub}):
            _doc_hash, records, _warnings = _document_chunks(
                "procedure.pdf", b"mock pdf", "tenant_docs"
            )

        self.assertEqual(records[0]["tenant_id"], "tenant_docs")
        self.assertEqual(records[0]["page_number"], 1)

    def test_powerpoint_parser_keeps_slide_title_text_and_tables(self):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Deployment policy"
        slide.placeholders[1].text = "Pause rollout when errors increase."
        table_shape = slide.shapes.add_table(
            2, 2, Inches(1), Inches(3), Inches(5), Inches(1)
        )
        table_shape.table.cell(0, 0).text = "Signal"
        table_shape.table.cell(0, 1).text = "Action"
        table_shape.table.cell(1, 0).text = "Errors"
        table_shape.table.cell(1, 1).text = "Pause"
        stream = BytesIO()
        presentation.save(stream)

        sections, warnings = _extract_sections("deployment.pptx", stream.getvalue())

        self.assertEqual(warnings, [])
        self.assertEqual(sections[0]["slide_number"], 1)
        self.assertEqual(sections[0]["section"], "Deployment policy")
        self.assertIn("Pause rollout when errors increase.", sections[0]["text"])
        self.assertIn("Errors | Pause", sections[0]["text"])


if __name__ == "__main__":
    unittest.main()
